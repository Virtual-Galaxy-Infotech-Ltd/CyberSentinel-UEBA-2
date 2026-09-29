"""Confirm every console workspace and Overview panel is named in the deck.

The brief was 'mention every feature', which is a claim worth testing rather
than assuming — it is easy to drop one while trimming copy to fit."""
import sys
from pptx import Presentation

FEATURES = {
    # workspace tabs
    "Overview": ["Overview"],
    "Incidents": ["Incident"],
    "Alert Feed": ["Alert Feed", "alert feed"],
    "User Risk": ["User Risk"],
    "Campaigns": ["Campaign"],
    "Endpoints": ["Endpoint"],
    "Threat Map": ["Threat Map"],
    "False Positives": ["False Positive", "False-positive", "false-positive"],
    "Watch List": ["Watch List", "Watch lists", "watch list"],
    "Reports": ["Report"],
    # overview panels / charts
    "Behavioural Risk Meter": ["Behavioural Risk Meter"],
    "Security Activity Pattern": ["Security Activity Pattern"],
    "Alert Timeline": ["Alert Timeline"],
    "Recent Alerts": ["Recent alerts"],
    "Top Risk Users": ["top risk users"],
    "Score Distribution": ["Score Distribution"],
    "Hourly Heatmap": ["Hourly Alert Heatmap"],
    "Top Rule Descriptions": ["top rule descriptions", "Top Rule Descriptions"],
    "Recommended Actions": ["recommended actions"],
    "MITRE Coverage": ["MITRE ATT&CK Coverage"],
    "Engine Health": ["engine health"],
    # capabilities
    "Evidence": ["Evidence"],
    "Filtering": ["Filtering", "severity threshold"],
    "Timeline windows": ["timeline windows", "Seven timeline windows"],
    "Live streaming": ["streams live"],
    "Custom reports": ["Custom reports"],
    "Scheduling": ["Scheduled", "Hourly, daily, weekly"],
    "Export formats": ["export formats", "PDF, Excel, CSV, HTML, JSON and XML"],
    "SSO / roles": ["Single sign-on"],
    "Audit trail": ["Audit trail"],
    "Data sovereignty": ["outbound connection"],
    "Detection layers": ["Isolation Forest", "autoencoder", "HDBSCAN"],
}

prs = Presentation(sys.argv[1])
text = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    text.append((i, r.text))
blob = " ".join(t for _, t in text)

missing = []
for name, needles in FEATURES.items():
    hit = next((n for n in needles if n in blob), None)
    if hit:
        slide = next((i for i, t in text if hit in t), "?")
        print(f"  ok    {name:26} (slide {slide})")
    else:
        missing.append(name)
print()
if missing:
    print(f"  NOT MENTIONED ({len(missing)}): " + ", ".join(missing))
    sys.exit(1)
print(f"  all {len(FEATURES)} features mentioned across {len(prs.slides)} slides")
