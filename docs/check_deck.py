"""Flag text that would overflow its box — the usual failure mode of a
generated deck, and one you cannot see without rendering."""
import sys
from pptx import Presentation
from pptx.util import Pt

prs = Presentation(sys.argv[1])
SLIDE_W = prs.slide_width / 914400.0
problems = []
for idx, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        box_w_pt = (sh.width / 914400.0) * 72
        box_h_pt = (sh.height / 914400.0) * 72
        total = 0.0
        for p in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in p.runs)
            if not txt:
                continue
            size = next((r.font.size.pt for r in p.runs if r.font.size), 18.0)
            bold = any(r.font.bold for r in p.runs)
            # Arial averages ~0.50 em per char, a little wider when bold.
            cw = size * (0.545 if bold else 0.50)
            cpl = max(1, int(box_w_pt / cw))
            lines = max(1, -(-len(txt) // cpl))
            ls = p.line_spacing or 1.15
            total += lines * size * ls * 1.2
            total += (p.space_before.pt if p.space_before else 0)
            total += (p.space_after.pt if p.space_after else 0)
        if total > box_h_pt + 2:
            problems.append((idx, round(total), round(box_h_pt),
                             "".join(r.text for r in sh.text_frame.paragraphs[0].runs)[:60]))
print(f"  slides: {len(prs.slides)}")
if problems:
    print(f"  POSSIBLE OVERFLOW in {len(problems)} text box(es):")
    for s, need, have, sample in problems:
        print(f"    slide {s:2}: needs ~{need}pt in {have}pt  |  {sample!r}")
else:
    print("  no text box exceeds its height estimate")
