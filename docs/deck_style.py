#!/usr/bin/env python3
"""Shared slide furniture for the CyberSentinel UEBA decks.

Both decks — the short client deck and the in-depth product deck — are built
from these primitives so they stay visually identical and a change to the
styling lands in both. Arial throughout: Segoe UI is Windows-only and falls back
to a serif on macOS, which is how an earlier draft came to look wrong.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BG    = RGBColor(0x0A, 0x11, 0x1C)
PANEL = RGBColor(0x12, 0x1D, 0x2E)
CYAN  = RGBColor(0x00, 0xD4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM   = RGBColor(0x9F, 0xB2, 0xC4)
FAINT = RGBColor(0x64, 0x7A, 0x8E)
AMBER = RGBColor(0xFF, 0xB4, 0x5C)
GREEN = RGBColor(0x4A, 0xD9, 0x91)
FONT  = "Arial"

MARGIN = Inches(0.85)


class Deck:
    def __init__(self, footer="CyberSentinel UEBA"):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        self.CONTENT_W = self.W - 2 * MARGIN
        self.footer = footer
        self.n = 0

    # ── primitives ──────────────────────────────────────────────────────────
    def slide(self, bg=BG):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        return s

    def textbox(self, slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return tf

    def write(self, tf, text, size=18, color=WHITE, bold=False, italic=False,
              space_before=0, space_after=6, align=PP_ALIGN.LEFT, first=False,
              line_spacing=1.15):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = FONT
        return p

    def rule(self, slide, top, width=Inches(1.5), color=CYAN,
             height=Inches(0.045), left=MARGIN):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        bar.shadow.inherit = False
        return bar

    def panel(self, slide, left, top, width, height, fill=PANEL, line=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                    width, height)
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line:
            sh.line.color.rgb = line
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.adjustments[0] = 0.04
        return sh

    def chrome(self, slide, title, kicker=None):
        top = Inches(0.52)
        if kicker:
            tf = self.textbox(slide, MARGIN, top, self.CONTENT_W, Inches(0.3))
            self.write(tf, kicker.upper(), size=11.5, color=CYAN, bold=True,
                       first=True, space_after=0)
            top = top + Inches(0.34)
        tf = self.textbox(slide, MARGIN, top, self.CONTENT_W, Inches(0.72))
        self.write(tf, title, size=30, color=WHITE, bold=True, first=True,
                   space_after=0)
        self.rule(slide, top + Inches(0.72))
        self.n += 1
        ftf = self.textbox(slide, MARGIN, self.H - Inches(0.52),
                           self.CONTENT_W, Inches(0.28))
        self.write(ftf, self.footer, size=9.5, color=FAINT, first=True,
                   space_after=0)
        ntf = self.textbox(slide, self.W - MARGIN - Inches(1.0),
                           self.H - Inches(0.52), Inches(1.0), Inches(0.28))
        self.write(ntf, str(self.n), size=9.5, color=FAINT,
                   align=PP_ALIGN.RIGHT, first=True, space_after=0)
        return Inches(2.05)

    def _lead(self, s, top, lead):
        if not lead:
            return top
        tf = self.textbox(s, MARGIN, top - Inches(0.22), self.CONTENT_W, Inches(0.7))
        self.write(tf, lead, size=15, color=DIM, first=True, space_after=0)
        return top + Inches(0.52)

    # ── slide types ─────────────────────────────────────────────────────────
    def title_slide(self, title, subtitle, blurb, footnote):
        s = self.slide()
        self.rule(s, Inches(2.45), width=Inches(2.4), height=Inches(0.06))
        tf = self.textbox(s, MARGIN, Inches(2.75), Inches(10.9), Inches(2.7))
        self.write(tf, title, size=50, color=WHITE, bold=True, first=True,
                   space_after=8)
        self.write(tf, subtitle, size=24, color=CYAN, space_after=18)
        self.write(tf, blurb, size=15.5, color=DIM, space_after=0, line_spacing=1.3)
        tf2 = self.textbox(s, MARGIN, self.H - Inches(1.15), self.CONTENT_W,
                           Inches(0.6))
        self.write(tf2, footnote, size=11.5, color=FAINT, first=True, space_after=0)
        return s

    def divider(self, kicker, title, blurb=None):
        s = self.slide()
        self.rule(s, Inches(2.9), width=Inches(1.8), height=Inches(0.055))
        tf = self.textbox(s, MARGIN, Inches(3.2), Inches(10.6), Inches(2.0))
        self.write(tf, kicker.upper(), size=13, color=CYAN, bold=True, first=True,
                   space_after=10)
        self.write(tf, title, size=38, color=WHITE, bold=True, space_after=12)
        if blurb:
            self.write(tf, blurb, size=15, color=DIM, space_after=0, line_spacing=1.3)
        self.n += 1
        ntf = self.textbox(s, self.W - MARGIN - Inches(1.0), self.H - Inches(0.52),
                           Inches(1.0), Inches(0.28))
        self.write(ntf, str(self.n), size=9.5, color=FAINT, align=PP_ALIGN.RIGHT,
                   first=True, space_after=0)
        return s

    def bullets(self, title, kicker, items, lead=None, size=16.5):
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        tf = self.textbox(s, MARGIN, top, self.CONTENT_W, self.H - top - Inches(0.8))
        for i, it in enumerate(items):
            if isinstance(it, tuple):
                self.write(tf, it[0], size=size, color=CYAN, bold=True,
                           first=(i == 0), space_before=0 if i == 0 else 12,
                           space_after=2)
                self.write(tf, it[1], size=13.5, color=DIM, space_after=0,
                           line_spacing=1.18)
            else:
                self.write(tf, "—   " + it, size=size, color=WHITE,
                           first=(i == 0), space_before=0 if i == 0 else 10,
                           space_after=0)
        return s

    def cards(self, title, kicker, cards, lead=None, per_row=3, body=12):
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        gap = Inches(0.28)
        rows = (len(cards) + per_row - 1) // per_row
        cw = int((self.CONTENT_W - gap * (per_row - 1)) / per_row)
        avail = self.H - top - Inches(0.78)
        ch = int((avail - gap * (rows - 1)) / rows)
        for i, (head, txt) in enumerate(cards):
            r, c = divmod(i, per_row)
            left = MARGIN + c * (cw + gap)
            ctop = top + r * (ch + gap)
            self.panel(s, left, ctop, cw, ch)
            tf = self.textbox(s, left + Inches(0.26), ctop + Inches(0.22),
                              cw - Inches(0.52), ch - Inches(0.44))
            self.write(tf, head, size=15.5, color=CYAN, bold=True, first=True,
                       space_after=5)
            self.write(tf, txt, size=body, color=DIM, space_after=0,
                       line_spacing=1.18)
        return s

    def metrics(self, title, kicker, metrics, lead=None, footnote=None, per_row=4):
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        gap = Inches(0.26)
        per_row = min(per_row, len(metrics))
        cw = int((self.CONTENT_W - gap * (per_row - 1)) / per_row)
        ch = Inches(1.55)
        for i, (val, label) in enumerate(metrics):
            r, c = divmod(i, per_row)
            left = MARGIN + c * (cw + gap)
            ctop = top + r * (ch + gap)
            self.panel(s, left, ctop, cw, ch)
            tf = self.textbox(s, left + Inches(0.22), ctop + Inches(0.26),
                              cw - Inches(0.44), ch - Inches(0.5))
            self.write(tf, val, size=27, color=CYAN, bold=True, first=True,
                       space_after=3)
            self.write(tf, label, size=11.5, color=DIM, space_after=0,
                       line_spacing=1.1)
        if footnote:
            ftf = self.textbox(s, MARGIN, self.H - Inches(1.05), self.CONTENT_W,
                               Inches(0.5))
            self.write(ftf, footnote, size=11, color=FAINT, italic=True,
                       first=True, space_after=0)
        return s

    def pipeline(self, title, kicker, stages, lead=None):
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        n = len(stages)
        gap = Inches(0.16)
        cw = int((self.CONTENT_W - gap * (n - 1)) / n)
        ch = Inches(2.9)
        for i, (num, head, txt) in enumerate(stages):
            left = MARGIN + i * (cw + gap)
            self.panel(s, left, top, cw, ch)
            self.rule(s, top + Inches(0.34), width=Inches(0.42),
                      left=left + Inches(0.22))
            tf = self.textbox(s, left + Inches(0.22), top + Inches(0.52),
                              cw - Inches(0.44), ch - Inches(0.8))
            self.write(tf, num, size=12, color=FAINT, bold=True, first=True,
                       space_after=4)
            self.write(tf, head, size=14, color=WHITE, bold=True, space_after=5)
            self.write(tf, txt, size=11.5, color=DIM, space_after=0,
                       line_spacing=1.2)
        return s

    def spec(self, title, kicker, rows, lead=None, note=None, key_w=Inches(3.6)):
        """A parameter table: what it is, and what it is set to / why.

        Row heights are derived from how many lines each value actually needs,
        and the font steps down until the whole table fits the slide. A fixed
        row height silently clipped any value that wrapped to a second line —
        invisible until the file was opened, and the most common way a
        generated deck ships broken.
        """
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        avail = self.H - top - (Inches(1.05) if note else Inches(0.6))
        gap = Inches(0.09)
        val_w_pt = ((self.CONTENT_W - key_w - Inches(0.3)) / 914400.0) * 72
        key_w_pt = ((key_w - Inches(0.3)) / 914400.0) * 72

        size, heights = 13.0, []
        for size in (13.0, 12.5, 12.0, 11.5, 11.0, 10.5, 10.0):
            heights = []
            for k, v in rows:
                v_lines = max(1, -(-len(v) // max(1, int(val_w_pt / (size * 0.5)))))
                k_lines = max(1, -(-len(k) // max(1, int(key_w_pt / (size * 0.55)))))
                lines = max(v_lines, k_lines)
                h_pt = lines * size * 1.12 * 1.2 + 19
                heights.append(max(Inches(0.5), Inches(h_pt / 72.0)))
            if sum(heights, Inches(0)) + gap * (len(rows) - 1) <= avail:
                break

        t = top
        for (k, v), rh in zip(rows, heights):
            self.panel(s, MARGIN, t, self.CONTENT_W, rh)
            tf = self.textbox(s, MARGIN + Inches(0.28), t + Inches(0.14),
                              key_w - Inches(0.3), rh - Inches(0.28))
            self.write(tf, k, size=size + 0.5, color=CYAN, bold=True, first=True,
                       space_after=0, line_spacing=1.12)
            tf = self.textbox(s, MARGIN + key_w, t + Inches(0.15),
                              self.CONTENT_W - key_w - Inches(0.3),
                              rh - Inches(0.28))
            self.write(tf, v, size=size, color=DIM, first=True, space_after=0,
                       line_spacing=1.12)
            t = t + rh + gap
        if note:
            ntf = self.textbox(s, MARGIN, self.H - Inches(1.0), self.CONTENT_W,
                               Inches(0.5))
            self.write(ntf, note, size=11, color=FAINT, italic=True, first=True,
                       space_after=0)
        return s

    def split(self, title, kicker, left_head, left_items, right_head, right_items,
              lead=None, left_color=GREEN, right_color=AMBER, note=None):
        s = self.slide()
        top = self._lead(s, self.chrome(s, title, kicker), lead)
        cw = int((self.CONTENT_W - Inches(0.3)) / 2)
        h = self.H - top - Inches(1.0)
        for off, head, items, col in ((0, left_head, left_items, left_color),
                                      (cw + Inches(0.3), right_head, right_items,
                                       right_color)):
            self.panel(s, MARGIN + off, top, cw, h, line=col)
            tf = self.textbox(s, MARGIN + off + Inches(0.3), top + Inches(0.26),
                              cw - Inches(0.6), h - Inches(0.5))
            self.write(tf, head, size=16.5, color=col, bold=True, first=True,
                       space_after=9)
            for t in items:
                self.write(tf, "·   " + t, size=12, color=WHITE, space_after=6,
                           line_spacing=1.12)
        if note:
            ntf = self.textbox(s, MARGIN, self.H - Inches(0.95), self.CONTENT_W,
                               Inches(0.45))
            self.write(ntf, note, size=11, color=FAINT, italic=True, first=True,
                       space_after=0)
        return s

    def closing(self, title, blurb, footnote):
        s = self.slide()
        self.rule(s, Inches(2.6), width=Inches(2.4), height=Inches(0.06))
        tf = self.textbox(s, MARGIN, Inches(2.95), Inches(10.6), Inches(2.2))
        self.write(tf, title, size=44, color=WHITE, bold=True, first=True,
                   space_after=12)
        self.write(tf, blurb, size=16, color=DIM, space_after=0, line_spacing=1.35)
        tf2 = self.textbox(s, MARGIN, self.H - Inches(1.15), self.CONTENT_W,
                           Inches(0.6))
        self.write(tf2, footnote, size=11.5, color=FAINT, first=True, space_after=0)
        return s

    def save(self, path):
        self.prs.save(path)
        print(f"  written: {path}  ({len(self.prs.slides._sldIdLst)} slides)")
